"""Questions from an Excel (.xlsx) or CSV file laid out like the template: one question per row.

Each row is checked on its own and reported with a list of problem codes (empty = can be imported),
so the lecturer sees exactly which rows to fix. The web app turns the codes into Vietnamese text.
"""

import csv
import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = ["Câu hỏi", "Đáp án A", "Đáp án B", "Đáp án C", "Đáp án D", "Đáp án đúng", "Giải thích"]
EXAMPLES = [
    ["She ___ to school every day.", "go", "goes", "gone", "going", "B", "Chủ ngữ số ít ở thì hiện tại đơn: động từ thêm -es."],
    ["They ___ football yesterday.", "play", "plays", "played", "playing", "C", "Có yesterday nên dùng thì quá khứ đơn."],
]
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_BYTES = 1_000_000
MAX_QUESTIONS = 200
#: Same limits as QuestionIn
MAX_LENGTH = {"content": 5000, "option": 1000, "explanation": 5000}


class ImportFileError(Exception):
    """The file itself cannot be used (type, size, unreadable, no rows, too many rows)."""


@dataclass
class ImportRow:
    row: int
    content: str
    option_a: str
    option_b: str
    option_c: str
    option_d: str
    correct_option: str | None
    explanation: str | None
    problems: list[str] = field(default_factory=list)


def template() -> bytes:
    """The workbook lecturers fill in: column names, two examples, A to D picker for the correct answer."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Câu hỏi"
    sheet.append(COLUMNS)
    for example in EXAMPLES:
        sheet.append(example)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0066CA")
        cell.alignment = Alignment(vertical="center")
    for letter, width in zip("ABCDEFG", (48, 18, 18, 18, 18, 14, 48)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    picker = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=False)
    picker.add(f"F2:F{MAX_QUESTIONS + 1}")
    sheet.add_data_validation(picker)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_rows(filename: str, data: bytes) -> list[tuple[int, list[str]]]:
    """(row number in the file, cells as text) for every non-empty row under the header row."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        rows = _xlsx_rows(data)
    elif name.endswith(".csv"):
        rows = _csv_rows(data)
    else:
        raise ImportFileError("Only .xlsx and .csv files are supported")
    # Row 1 holds the column names; numbers match what the lecturer sees in Excel
    numbered = [(index, cells) for index, cells in enumerate(rows, start=1) if index > 1 and any(cells)]
    if not numbered:
        raise ImportFileError("The file has no questions")
    if len(numbered) > MAX_QUESTIONS:
        raise ImportFileError(f"At most {MAX_QUESTIONS} questions per file")
    return numbered


def check(rows: list[tuple[int, list[str]]]) -> list[ImportRow]:
    """Turns the cells into questions and lists what is wrong with each one."""
    checked = []
    for number, cells in rows:
        content, option_a, option_b, option_c, option_d, correct, explanation = (cells + [""] * 7)[:7]
        problems = []
        if not content:
            problems.append("content_missing")
        for letter, text in zip("abcd", (option_a, option_b, option_c, option_d)):
            if not text:
                problems.append(f"option_{letter}_missing")
        correct = correct.upper()
        if correct not in ("A", "B", "C", "D"):
            problems.append("correct_option_invalid")
        too_long = (
            len(content) > MAX_LENGTH["content"]
            or any(len(text) > MAX_LENGTH["option"] for text in (option_a, option_b, option_c, option_d))
            or len(explanation) > MAX_LENGTH["explanation"]
        )
        if too_long:
            problems.append("too_long")
        checked.append(
            ImportRow(
                row=number,
                content=content,
                option_a=option_a,
                option_b=option_b,
                option_c=option_c,
                option_d=option_d,
                correct_option=correct if correct in ("A", "B", "C", "D") else None,
                explanation=explanation or None,
                problems=problems,
            )
        )
    return checked


def same_question_key(content: str) -> str:
    """Spaces and letter case do not make a question different."""
    return " ".join(content.split()).lower()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 1990 as 1990.0
    return str(value).strip()


def _xlsx_rows(data: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # not a zip, not a workbook, damaged file...
        raise ImportFileError("The file could not be read as an Excel workbook") from exc
    try:
        return [[_text(value) for value in row] for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _csv_rows(data: bytes) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")  # Excel adds a BOM when saving "CSV UTF-8"
    except UnicodeDecodeError as exc:
        raise ImportFileError("CSV files must be saved as UTF-8") from exc
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text))]
