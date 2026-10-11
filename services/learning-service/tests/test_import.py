"""Phase 6: lecturers import questions from an Excel or CSV file laid out like the template."""

import io

import pytest
from openpyxl import Workbook, load_workbook

from learning_service import question_import

TEACHER = ("3", "TEACHER")
STUDENT = ("2", "STUDENT")


@pytest.fixture
def h(auth_header):
    return lambda who: auth_header(*who)


@pytest.fixture
def topic(bank, h):
    """A topic that already has one question."""
    topic = bank.post("/topics", json={"name": "Thì hiện tại đơn"}, headers=h(TEACHER)).json()
    body = {
        "topicId": topic["id"],
        "content": "She ___ to school every day.",
        "optionA": "go",
        "optionB": "goes",
        "optionC": "gone",
        "optionD": "going",
        "correctOption": "B",
    }
    assert bank.post("/questions", json=body, headers=h(TEACHER)).status_code == 201
    return topic


def workbook(*rows) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(question_import.COLUMNS)
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(client, h, topic_id, data, filename="cau-hoi.xlsx", dry_run=True, who=TEACHER):
    return client.post(
        f"/questions/import?topicId={topic_id}&dryRun={'true' if dry_run else 'false'}",
        files={"file": (filename, data, "application/octet-stream")},
        headers=h(who),
    )


ROWS = (
    ("Q1: He ___ tea every morning.", "drink", "drinks", "drank", "drinking", "b", "Chủ ngữ số ít."),  # row 2, lower-case answer is fine
    ("Q2: missing C", "a", "b", None, "d", "A", None),  # row 3
    ("Q3: wrong answer letter", "a", "b", "c", "d", "E", None),  # row 4
    (None, None, None, None, None, None, None),  # row 5, empty: skipped
    ("she ___ to school   every day.", "go", "goes", "gone", "going", "B", None),  # row 6, already in the topic
    ("Q1: He ___ tea every morning.", "drink", "drinks", "drank", "drinking", "B", None),  # row 7, repeated in the file
    ("1 + 1 = ?", 1, 2, 3, 4, "B", None),  # row 8, numbers become text
)


def test_template_has_the_expected_columns_and_examples(bank, h):
    response = bank.get("/questions/import/template", headers=h(TEACHER))
    assert response.status_code == 200
    assert response.headers["content-type"] == question_import.XLSX_TYPE
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == question_import.COLUMNS
    assert len(rows) == 3
    assert bank.get("/questions/import/template", headers=h(STUDENT)).status_code == 403


def test_preview_reports_each_row_and_saves_nothing(bank, h, topic):
    response = upload(bank, h, topic["id"], workbook(*ROWS))
    assert response.status_code == 200
    body = response.json()
    assert (body["dryRun"], body["total"], body["valid"], body["imported"]) == (True, 6, 2, 0)
    assert [(row["row"], row["problems"]) for row in body["rows"]] == [
        (2, []),
        (3, ["option_c_missing"]),
        (4, ["correct_option_invalid"]),
        (6, ["duplicate"]),
        (7, ["duplicate"]),
        (8, []),
    ]
    assert body["rows"][0]["correctOption"] == "B"
    assert (body["rows"][5]["optionA"], body["rows"][5]["optionD"]) == ("1", "4")
    assert len(bank.get(f"/questions?topicId={topic['id']}", headers=h(TEACHER)).json()) == 1


def test_import_saves_the_valid_rows_into_the_topic(bank, h, topic):
    response = upload(bank, h, topic["id"], workbook(*ROWS), dry_run=False)
    assert response.status_code == 200
    assert (response.json()["valid"], response.json()["imported"]) == (2, 2)
    questions = bank.get(f"/questions?topicId={topic['id']}", headers=h(TEACHER)).json()
    assert len(questions) == 3
    imported = next(question for question in questions if question["content"].startswith("Q1"))
    assert (imported["correctOption"], imported["explanation"], imported["createdBy"]) == ("B", "Chủ ngữ số ít.", 3)
    # Importing the same file again adds nothing: every row is now a duplicate or broken
    again = upload(bank, h, topic["id"], workbook(*ROWS), dry_run=False).json()
    assert (again["valid"], again["imported"]) == (0, 0)


def test_a_broken_row_does_not_block_a_correct_copy_later(bank, h, topic):
    rows = (("New question", "a", "b", None, "d", "A", None), ("New question", "a", "b", "c", "d", "A", None))
    body = upload(bank, h, topic["id"], workbook(*rows)).json()
    assert [row["problems"] for row in body["rows"]] == [["option_c_missing"], []]


def test_csv_files_saved_by_excel_work_too(bank, h, topic):
    text = "Câu hỏi,Đáp án A,Đáp án B,Đáp án C,Đáp án D,Đáp án đúng,Giải thích\nThey ___ football yesterday.,play,plays,played,playing,C,Quá khứ đơn.\n"
    response = upload(bank, h, topic["id"], text.encode("utf-8-sig"), filename="cau-hoi.csv", dry_run=False)
    assert response.status_code == 200
    assert response.json()["imported"] == 1
    assert upload(bank, h, topic["id"], "Câu hỏi\nđ".encode("utf-16"), filename="cu.csv").status_code == 400


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("cau-hoi.txt", b"hello", "Only .xlsx and .csv files are supported"),
        ("cau-hoi.xlsx", b"not a workbook", "The file could not be read as an Excel workbook"),
        ("cau-hoi.xlsx", workbook(), "The file has no questions"),
        ("cau-hoi.xlsx", workbook(*[(f"Q{i}", "a", "b", "c", "d", "A", None) for i in range(201)]), "At most 200 questions per file"),
    ],
)
def test_unusable_files_are_refused(bank, h, topic, filename, data, message):
    response = upload(bank, h, topic["id"], data, filename=filename)
    assert response.status_code == 400
    assert response.json()["message"] == message


def test_import_needs_an_existing_topic_and_a_lecturer(bank, h, topic):
    assert upload(bank, h, 999, workbook(*ROWS)).status_code == 404
    assert upload(bank, h, topic["id"], workbook(*ROWS), who=STUDENT).status_code == 403
    too_big = upload(bank, h, topic["id"], b"x" * (question_import.MAX_BYTES + 1), filename="big.csv")
    assert (too_big.status_code, too_big.json()["message"]) == (400, "The file is larger than 1 MB")
